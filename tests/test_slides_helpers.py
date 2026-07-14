import slides_export as sx


def test_new_deck_is_16x9():
    prs = sx.new_deck()
    # 13.333in x 7.5in in EMU (914400 EMU/in)
    assert abs(prs.slide_width - 12192000) < 2000
    assert abs(prs.slide_height - 6858000) < 2000


def test_blank_slide_and_helpers_run():
    prs = sx.new_deck()
    slide = sx.blank_slide(prs)
    sx.rect(slide, 0, 0, 13.333, 7.5, fill=sx.WHITE)
    sx.text(slide, 0.5, 0.5, 6, 0.5, ("Hello", 24, sx.INK, sx.F_BOLD, True, None))
    sx.stat_tile(slide, 0.5, 1.5, 2.9, 1.2, "REVENUE", "SAR 1.2K", "▲ 5%", sx.GREEN)
    sx.styled_table(slide, 0.5, 3.0, 5.0, ["A", "B"], [("x", "1"), ("y", "2")])
    sx.insight_box(slide, 7.0, 3.0, 5.0, 1.0, "KEY INSIGHT", "Body text.")
    sx.recommendation_strip(slide, 6.4, "Do the thing.", "Journey Designer")
    sx.add_morph(slide)
    # morph transition present in slide xml
    from pptx.oxml.ns import qn
    assert slide.element.find(qn("p:transition")) is not None


def test_fig_to_png_returns_png_bytes():
    import plotly.graph_objects as go
    png = sx.fig_to_png(go.Figure(go.Bar(x=[1], y=[1])))
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_channel_status_thresholds():
    assert sx.channel_status(0)[0] == "INACTIVE"
    assert sx.channel_status(50)[0] == "LOW VOLUME"
    assert sx.channel_status(500)[0] == "ACTIVE"


def test_qoq_delta_text_direction():
    up_text, up_color = sx.qoq_delta_text({"pct_change": 12.5}, "prior period")
    assert "12.5%" in up_text and up_color == sx.GREEN
    down_text, down_color = sx.qoq_delta_text({"pct_change": -8.0}, "prior period")
    assert "-8.0%" in down_text and down_color == sx.RED


def test_status_pill_draws_without_error():
    prs = sx.new_deck()
    slide = sx.blank_slide(prs)
    width = sx.status_pill(slide, 0.5, 0.5, "ACTIVE", sx.GREEN)
    assert width > 0


def test_header_renders_story_line_when_given():
    prs = sx.new_deck()
    slide = sx.blank_slide(prs)
    sx._header(slide, "TREND", "Conversions Over Time", "Jun 2026", story="Here's the shape behind that total.")
    texts = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame) if False else \
            " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    assert "Here's the shape behind that total." in texts
